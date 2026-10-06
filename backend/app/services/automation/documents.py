"""Documents the assistant may attach to a file field (V0.10).

Only files that belong to **this user's application** can be chosen, and the choice
is made by *id* -- a generated version of this application's CV or Anschreiben, or
the resume file the application is built on. The server resolves it; no client
ever supplies a path, and nothing here exposes one.

A generated document is stored as structured content, so it is rendered to a DOCX
in memory when it is attached (nothing is written to disk). A choice is pinned to
one specific immutable version: a newer version never replaces it by itself.
"""

import io
import mimetypes
import re
import uuid
from dataclasses import dataclass

from docx import Document
from docx.document import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.logging import get_logger
from app.models.application import Application
from app.models.generated_document import (
    GeneratedDocument,
    GeneratedDocumentOrigin,
    GeneratedDocumentStatus,
    GeneratedDocumentType,
)
from app.models.resume import Resume
from app.models.user import User
from app.schemas.document_content import CONTENT_ADAPTER, CoverLetterContent, TailoredCvContent
from app.services.storage.factory import get_file_storage

logger = get_logger(__name__)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024

_TYPE_FOR = {
    "tailored_cv": GeneratedDocumentType.TAILORED_CV,
    "cover_letter": GeneratedDocumentType.COVER_LETTER,
}
_ORIGIN_SUFFIX = {
    GeneratedDocumentOrigin.AI: "",
    GeneratedDocumentOrigin.HUMANIZED: " (reworded by Humanize)",
    GeneratedDocumentOrigin.USER_EDITED: " (edited by you)",
}


class DocumentUnavailable(Exception):
    """The chosen document cannot be attached any more."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class DocumentChoice:
    type: str  # tailored_cv | cover_letter | resume_file
    id: uuid.UUID
    label: str
    version_number: int | None
    origin: str
    is_current: bool
    file_name: str
    mime_type: str


def _safe_name(text: str) -> str:
    cleaned = re.sub(r"[^\w .()-]+", "", text, flags=re.UNICODE).strip(" .")
    return re.sub(r"\s+", " ", cleaned)[:80] or "Document"


def _generated_choice(row: GeneratedDocument, *, current: int | None) -> DocumentChoice:
    german = row.language == "de"
    kind = row.document_type
    title = (
        ("Tailored CV" if not german else "Lebenslauf")
        if kind is GeneratedDocumentType.TAILORED_CV
        else ("Cover letter" if not german else "Anschreiben")
    )
    person = _person_of(row)
    file_name = _safe_name(f"{person + ' - ' if person else ''}{title}") + ".docx"
    english_title = "Tailored CV" if kind is GeneratedDocumentType.TAILORED_CV else "Cover letter"
    return DocumentChoice(
        type="tailored_cv" if kind is GeneratedDocumentType.TAILORED_CV else "cover_letter",
        id=row.id,
        label=f"{english_title} v{row.version_number}{_ORIGIN_SUFFIX[row.origin]}",
        version_number=row.version_number,
        origin=row.origin.value,
        is_current=row.version_number == current,
        file_name=file_name,
        mime_type=DOCX_MIME,
    )


def _person_of(row: GeneratedDocument) -> str | None:
    content = row.content or {}
    header = content.get("header") or content.get("sender") or {}
    name = header.get("full_name")
    return str(name) if name else None


def _resume_choice(resume: Resume) -> DocumentChoice:
    return DocumentChoice(
        type="resume_file",
        id=resume.id,
        label=f"Your uploaded resume: {resume.file_name}",
        version_number=resume.version_number,
        origin="upload",
        is_current=True,
        file_name=resume.file_name,
        mime_type=resume.mime_type,
    )


def _validated_documents(
    db: Session, user: User, application: Application
) -> list[GeneratedDocument]:
    return list(
        db.scalars(
            select(GeneratedDocument)
            .where(
                GeneratedDocument.application_id == application.id,
                GeneratedDocument.user_id == user.id,
                GeneratedDocument.status == GeneratedDocumentStatus.VALIDATED,
            )
            .order_by(GeneratedDocument.version_number.desc())
        )
    )


def _application_resume(user: User, application: Application) -> Resume | None:
    resume = application.resume
    return resume if resume is not None and resume.user_id == user.id else None


def list_choices(db: Session, user: User, application: Application) -> list[DocumentChoice]:
    """Everything the user may attach for this application, newest version first."""
    rows = _validated_documents(db, user, application)
    current: dict[GeneratedDocumentType, int] = {}
    for row in rows:
        if row.version_number is not None:
            current[row.document_type] = max(current.get(row.document_type, 0), row.version_number)
    choices = [_generated_choice(row, current=current.get(row.document_type)) for row in rows]
    resume = _application_resume(user, application)
    if resume is not None:
        choices.append(_resume_choice(resume))
    return choices


def find_choice(
    db: Session,
    user: User,
    application: Application,
    document_type: str,
    document_id: uuid.UUID,
) -> DocumentChoice | None:
    """The choice with that type and id -- only if it is this user's, for this application."""
    for choice in list_choices(db, user, application):
        if choice.type == document_type and choice.id == document_id:
            return choice
    return None


def proposed_choice(
    db: Session, user: User, application: Application, detected_type: str | None
) -> DocumentChoice | None:
    """What "Use this file" would pick: the current version of the document type the
    field clearly asks for. A field that does not say gets no proposal."""
    if detected_type is None:
        return None
    candidates = [c for c in list_choices(db, user, application) if c.type == detected_type]
    return next((c for c in candidates if c.is_current), None)


# --- Does the form take this file? ---------------------------------------------------------------


def file_accepted(accept: str | None, file_name: str, mime_type: str) -> bool:
    """Whether a file input's `accept` list allows the file. No list: anything."""
    if not accept or not accept.strip():
        return True
    extension = "." + file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    for token in (t.strip().lower() for t in accept.split(",")):
        if not token:
            continue
        if token.startswith(".") and token == extension:
            return True
        if token == mime_type.lower():
            return True
        if token.endswith("/*") and mime_type.lower().startswith(token[:-1]):
            return True
        if token == "*/*":
            return True
    return False


# --- Rendering -----------------------------------------------------------------------------------


def _dates(entry: object) -> str:
    text = getattr(entry, "date_text", None)
    if text:
        return str(text)
    start = getattr(entry, "start_date", None) or ""
    end = getattr(entry, "end_date", None) or (
        "present" if getattr(entry, "is_current", False) else ""
    )
    return f"{start} – {end}".strip(" –") if start or end else ""


def _cv_docx(content: TailoredCvContent) -> DocxDocument:
    doc = Document()
    header = content.header
    if header.full_name:
        doc.add_heading(header.full_name, level=0)
    contact = [header.email, header.phone, header.location, *header.links]
    contact_line = " · ".join(str(c) for c in contact if c)
    if contact_line:
        doc.add_paragraph(contact_line)
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(content.summary.text)
    if content.experience:
        doc.add_heading("Experience", level=1)
        for entry in content.experience:
            doc.add_paragraph(
                " — ".join(p for p in (entry.job_title, entry.company) if p), style="Heading 2"
            )
            meta = " · ".join(p for p in (entry.location, _dates(entry)) if p)
            if meta:
                doc.add_paragraph(meta)
            for bullet in entry.bullets:
                doc.add_paragraph(bullet.text, style="List Bullet")
    if content.skills:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(skill.name for skill in content.skills))
    if content.education:
        doc.add_heading("Education", level=1)
        for edu in content.education:
            line = ", ".join(p for p in (edu.degree, edu.field_of_study, edu.institution) if p)
            dates = _dates(edu)
            doc.add_paragraph(f"{line} ({dates})" if dates else line)
    if content.languages:
        doc.add_heading("Languages", level=1)
        for lang in content.languages:
            doc.add_paragraph(
                f"{lang.language}{' — ' + lang.proficiency if lang.proficiency else ''}"
            )
    if content.certifications:
        doc.add_heading("Certifications", level=1)
        for cert in content.certifications:
            doc.add_paragraph(", ".join(p for p in (cert.name, cert.issuer, cert.date) if p))
    if content.projects:
        doc.add_heading("Projects", level=1)
        for project in content.projects:
            doc.add_paragraph(
                f"{project.name}{' (' + project.role + ')' if project.role else ''}"
                f"{' — ' + project.description if project.description else ''}"
            )
    if content.additional_information:
        doc.add_heading("Additional information", level=1)
        for info in content.additional_information:
            doc.add_paragraph(info.value)
    return doc


def _letter_docx(content: CoverLetterContent) -> DocxDocument:
    doc = Document()
    sender = [
        content.sender.full_name,
        content.sender.email,
        content.sender.phone,
        content.sender.location,
    ]
    for line in (s for s in sender if s):
        doc.add_paragraph(line)
    if content.recipient.company_name:
        doc.add_paragraph(content.recipient.company_name)
    doc.add_paragraph(content.subject).runs[0].bold = True
    doc.add_paragraph(content.salutation)
    for paragraph in content.paragraphs:
        doc.add_paragraph(paragraph.text)
    if content.salary_statement:
        doc.add_paragraph(content.salary_statement)
    closing = [p for p in (content.closing_formula, content.signature) if p]
    if closing:
        doc.add_paragraph("\n\n".join(closing))
    return doc


def render(
    db: Session, user: User, application: Application, choice: DocumentChoice
) -> tuple[str, str, bytes]:
    """(file name, MIME type, bytes) of the chosen document, read by the server.
    Raises `DocumentUnavailable` -- never leaks a path."""
    # Always re-resolve: the choice must still be this user's, for this application.
    live = find_choice(db, user, application, choice.type, choice.id)
    if live is None:
        raise DocumentUnavailable("DOCUMENT_UNAVAILABLE")
    if live.type == "resume_file":
        resume = _application_resume(user, application)
        if resume is None or resume.id != live.id:
            raise DocumentUnavailable("DOCUMENT_UNAVAILABLE")
        try:
            data = get_file_storage().read(reference=resume.storage_reference)
        except FileNotFoundError:
            raise DocumentUnavailable("DOCUMENT_UNAVAILABLE") from None
        mime = (
            resume.mime_type
            or mimetypes.guess_type(resume.file_name)[0]
            or "application/octet-stream"
        )
        return resume.file_name, mime, data
    row = db.scalar(
        select(GeneratedDocument).where(
            GeneratedDocument.id == live.id,
            GeneratedDocument.application_id == application.id,
            GeneratedDocument.user_id == user.id,
            GeneratedDocument.status == GeneratedDocumentStatus.VALIDATED,
        )
    )
    if row is None or row.content is None:
        raise DocumentUnavailable("DOCUMENT_UNAVAILABLE")
    content = CONTENT_ADAPTER.validate_python(row.content)
    doc = _cv_docx(content) if isinstance(content, TailoredCvContent) else _letter_docx(content)
    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()
    if len(data) > MAX_ATTACHMENT_BYTES:
        raise DocumentUnavailable("DOCUMENT_UNAVAILABLE")
    return live.file_name, DOCX_MIME, data
