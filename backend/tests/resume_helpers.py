"""Deterministic helpers for Resume Intelligence tests.

No test ever calls a real (paid) AI API: `FakeAIProvider` returns canned
responses from fixtures and validates them exactly like a real provider.
"""

import io
import json
from collections.abc import Callable
from pathlib import Path
from typing import Any

from pydantic import ValidationError

from app.ai.provider import AIInvalidResponseError, AIProvider, T, summarize_validation_error

FIXTURES = Path(__file__).parent / "fixtures" / "resume"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def sample_resume_text() -> str:
    return (FIXTURES / "sample_resume.txt").read_text(encoding="utf-8")


def expected_parsed_resume() -> dict[str, Any]:
    return json.loads((FIXTURES / "expected_parsed_resume.json").read_text(encoding="utf-8"))


class FakeAIProvider(AIProvider):
    """Returns a canned structured response (or raises a canned error)."""

    name = "fake"
    model = "fake-model-1"

    def __init__(
        self,
        response: dict[str, Any] | Exception | None = None,
        on_call: Callable[[], None] | None = None,
    ) -> None:
        self.response: dict[str, Any] | Exception = (
            response if response is not None else expected_parsed_resume()
        )
        self.on_call = on_call
        self.calls: list[dict[str, str]] = []

    async def generate_structured(self, *, system: str, prompt: str, schema: type[T]) -> T:
        self.calls.append({"system": system, "prompt": prompt})
        if self.on_call is not None:
            self.on_call()
        if isinstance(self.response, Exception):
            raise self.response
        try:
            return schema.model_validate(self.response)
        except ValidationError as exc:
            raise AIInvalidResponseError(summarize_validation_error(exc)) from None


def _pdf_escape(text: str) -> bytes:
    encoded = text.encode("cp1252")
    return encoded.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def make_text_pdf(text: str) -> bytes:
    """A minimal, valid single-page PDF whose text layer contains `text`."""
    lines = text.split("\n")
    stream = b"BT /F1 9 Tf 11 TL 40 800 Td\n"
    for line in lines:
        stream += b"(" + _pdf_escape(line) + b") Tj T*\n"
    stream += b"ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    return _assemble_pdf(objects)


def _assemble_pdf(objects: list[bytes]) -> bytes:
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_blank_pdf(pages: int = 1) -> bytes:
    """A PDF with pages but no text layer -- what a scanned resume looks like."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=842)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(
    text: str, *, table_rows: list[list[str]] | None = None, header: str | None = None
) -> bytes:
    from docx import Document

    document = Document()
    if header:
        document.sections[0].header.paragraphs[0].text = header
    for line in text.split("\n"):
        document.add_paragraph(line)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
